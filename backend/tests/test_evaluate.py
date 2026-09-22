"""Decisions, totals and export.

The numbers here are the answer key from CLAUDE.md, verified by hand against
the supplied files. If a change breaks one, the change is wrong until proven
otherwise.
"""

import io
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.evaluate import (
    Change,
    EvaluatedOffer,
    evaluate,
    prepare_change,
    store,
    validate_change,
)
from app.core.export import (
    COLUMNS,
    escape_text,
    export_records,
    export_rows,
    left_out_records,
    to_csv,
)
from app.core.export_xlsx import (
    COLUMN_SPEC,
    SHEET_NAME,
    SUMMARY_SHEET,
    WorkbookContext,
    read_amount,
    to_xlsx,
)
from app.core.money import format_amount, to_units
from app.core.normalize import Parsed, parse_money, parse_quantity, parse_text
from app.core.reader import SourceCell, SourceLine, read_workbook
from app.core.rules import analyze

FIXTURES = Path(__file__).parent / "fixtures"

ALL_FIXTURES = [
    "01-northstar-line-sheet.xlsx",
    "02-harbor-size-grid.xlsx",
    "03-northstar-5000-rows.xlsx",
]


def stored(name: str):
    result = read_workbook((FIXTURES / name).read_bytes())
    return result, store(analyze(result.lines))


def save(lines, decisions, *changes):
    """Apply changes the way the API does: validate everything, then all or nothing."""
    by_line = {line.line_id: line for line in lines}
    prepared, errors = [], {}
    for change in changes:
        decision, problems = prepare_change(by_line.get(change.line_id), change)
        if problems:
            errors[change.line_id] = problems
        prepared.append((change, decision))
    assert not errors, errors

    state = {d.line_id: d for d in decisions}
    for change, decision in prepared:
        if change.action == "reset":
            state.pop(change.line_id, None)
        else:
            state[change.line_id] = decision
    return list(state.values())


def synthetic(*specs, description="Tee", retail="10.00"):
    """StoredLines built by hand, for rules that no fixture happens to hit."""
    sources = []
    for index, (code, size, qty, cost) in enumerate(specs, start=1):
        row = index + 1
        fields = {"item_code": code, "description": description, "size": size,
                  "category": None, "quantity": qty, "unit_cost": cost, "retail": retail}
        sources.append(SourceLine(
            line_id=f"L{index}", source_row=row,
            item_code=parse_text(code), description=parse_text(description),
            size=parse_text(size), category=Parsed(None, "missing"),
            quantity=parse_quantity(qty), unit_cost=parse_money(cost),
            retail=parse_money(retail),
            cells={n: SourceCell(f"A{row}", v) for n, v in fields.items()},
        ))
    return store(analyze(sources))


def issue_codes(line):
    return {i.code for i in line.issues}


# ---------- the answer key survives evaluate() ----------

@pytest.mark.parametrize("name, lines, pieces, cost, retail", [
    ("01-northstar-line-sheet.xlsx", 6, 1_733, "4208.00", "21552.00"),
    ("02-harbor-size-grid.xlsx", 22, 1_340, "3740.00", "16144.00"),
    ("03-northstar-5000-rows.xlsx", 5_000, 62_444, "187214.50", "1031508.00"),
])
def test_default_totals_match_the_answer_key(name, lines, pieces, cost, retail):
    _, stored_lines = stored(name)
    summary = evaluate(stored_lines, []).summary
    assert summary.included_lines == lines
    assert summary.pieces == pieces
    assert format_amount(summary.supplier_cost_units) == cost
    assert format_amount(summary.retail_reference_units) == retail


def test_retail_never_enters_the_supplier_cost_total():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    from_cost_alone = sum(
        line.quantity * line.unit_cost_units for line in offer.lines if line.included
    )
    assert offer.summary.supplier_cost_units == from_cost_alone
    assert offer.summary.retail_reference_units != from_cost_alone


# ---------- decisions ----------

def test_including_the_conflicted_a108_row_at_five_dollars():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, save(lines, [], Change("R14", "include", unit_cost="5.00")))

    assert offer.summary.pieces == 1_833
    assert format_amount(offer.summary.supplier_cost_units) == "4708.00"
    # The other half of the conflict stays out; choosing one is not choosing both.
    assert offer.by_id()["R15"].status == "excluded"


def test_fixing_a104s_missing_cost_adds_its_pieces_and_value():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    base = evaluate(lines, []).summary
    offer = evaluate(lines, save(lines, [], Change("R9", "include", unit_cost="9.00")))

    assert offer.summary.pieces == base.pieces + 60
    assert offer.summary.supplier_cost_units == base.supplier_cost_units + to_units(Decimal("540"))
    line = offer.by_id()["R9"]
    assert line.included and line.unit_cost_units == to_units(Decimal("9.00"))
    assert line.overridden_fields == ("unit_cost_units",)
    # The original parse stays visible next to the effective value.
    assert line.original.unit_cost_units.status == "missing"


def test_excluding_a_clean_line_removes_exactly_that_line():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    before = evaluate(lines, [])
    target = next(line for line in before.lines if line.included and not line.issues)

    after = evaluate(lines, save(lines, [], Change(target.line_id, "exclude", note="Sold already")))

    assert after.summary.included_lines == before.summary.included_lines - 1
    assert after.summary.pieces == before.summary.pieces - target.quantity
    assert after.summary.supplier_cost_units == (
        before.summary.supplier_cost_units - target.line_value_units
    )
    assert after.summary.user_excluded == 1
    assert after.by_id()[target.line_id].status_reason == "You excluded this line. Sold already"


def test_reset_restores_the_default():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    decisions = save(lines, [], Change("R14", "include", unit_cost="5.00"))
    decisions = save(lines, decisions, Change("R14", "reset"))

    assert decisions == []
    summary = evaluate(lines, decisions).summary
    assert summary.pieces == 1_733
    assert format_amount(summary.supplier_cost_units) == "4208.00"


def test_a_blank_size_can_be_fixed_instead_of_being_a_dead_end():
    line = synthetic(("A100", None, 10, "3.00"))[0]
    assert "MISSING_SIZE" in {i.code for i in line.issues}
    assert any("no size" in e for e in validate_change(line, Change("L1", "include")))

    decision, errors = prepare_change(line, Change("L1", "include", size="M"))
    assert errors == []
    included = evaluate([line], [decision]).by_id()["L1"]
    assert included.included and included.size == "M"
    assert included.original.size.status == "missing"


# ---------- a duplicate created by an edit ----------

def test_including_both_halves_of_a_conflict_is_warned_not_blocked():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, save(lines, [], Change("R14", "include"), Change("R15", "include")))

    by_line = offer.by_id()
    for line_id in ("R14", "R15"):
        assert by_line[line_id].included
        assert "DUPLICATE_AFTER_EDIT" in issue_codes(by_line[line_id])


def test_an_edit_that_creates_a_duplicate_warns_but_never_blocks():
    lines = synthetic(("A100", "M", 10, "3.00"), ("A200", "M", 5, "4.00"))
    offer = evaluate(lines, save(lines, [], Change("L2", "include", item_code="A100")))

    by_line = offer.by_id()
    assert by_line["L1"].included and by_line["L2"].included
    assert all("DUPLICATE_AFTER_EDIT" in issue_codes(by_line[i]) for i in ("L1", "L2"))
    assert offer.summary.pieces == 15
    assert offer.summary.warnings == 2


def test_no_duplicate_warning_when_nothing_collides():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    assert not any("DUPLICATE_AFTER_EDIT" in issue_codes(line) for line in offer.lines)


# ---------- validation, including straight past the UI ----------

@pytest.mark.parametrize("change, expected", [
    (Change("R9", "include", quantity=12.5), "whole number"),
    (Change("R9", "include", quantity="12"), "whole number"),
    (Change("R9", "include", quantity=True), "whole number"),
    (Change("R9", "include", quantity=12.0), "whole number"),
    (Change("R9", "include", quantity=0, unit_cost="9.00"), "between 1 and"),
    (Change("R9", "include", quantity=-5, unit_cost="9.00"), "between 1 and"),
    (Change("R9", "include", quantity=10_000_001, unit_cost="9.00"), "between 1 and"),
    (Change("R9", "include", unit_cost="abc"), "isn't an amount"),
    (Change("R9", "include", unit_cost="1.23456"), "isn't an amount"),
    (Change("R9", "include", unit_cost=9.0), "decimal string"),
    (Change("R9", "include", unit_cost=9), "decimal string"),
    (Change("R9", "include", unit_cost="-1"), "above $0"),
    (Change("R9", "include", unit_cost="0"), "above $0"),
    (Change("R9", "include", unit_cost="2000000"), "limit"),
    (Change("R9", "include", unit_cost="9.00", item_code="   "), "can't be blank"),
    (Change("R9", "include", unit_cost="9.00", item_code="x" * 65), "the most we store"),
    (Change("R9", "include", unit_cost="9.00", item_code=101), "must be text"),
    (Change("R9", "include", unit_cost="9.00", size="x" * 65), "the most we store"),
    (Change("R9", "include"), "no usable supplier cost"),
    (Change("R16", "include", unit_cost="1.00"), "no item code"),
    (Change("nope", "include"), "no line"),
    (Change("R9", "exclude", quantity=5), "can't be set while excluding"),
    (Change("R9", "reset", unit_cost="1.00"), "can't carry values"),
    (Change("R9", "sideways"), "Action must be"),
])
def test_invalid_changes_are_rejected_with_a_clear_message(change, expected):
    _, lines = stored("01-northstar-line-sheet.xlsx")
    by_line = {line.line_id: line for line in lines}

    errors = validate_change(by_line.get(change.line_id), change)

    assert errors, "expected this change to be rejected"
    assert any(expected in e for e in errors), errors


def test_a_rejected_change_produces_no_decision():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    by_line = {line.line_id: line for line in lines}
    decision, errors = prepare_change(by_line["R9"], Change("R9", "include", quantity="12"))
    assert decision is None and errors


def test_one_bad_field_does_not_produce_two_confusing_messages():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    by_line = {line.line_id: line for line in lines}
    errors = validate_change(by_line["R9"], Change("R9", "include", unit_cost="abc"))
    assert len(errors) == 1, errors


# ---------- export ----------

@pytest.mark.parametrize("name", ALL_FIXTURES)
def test_export_sums_equal_the_summary(name):
    result, lines = stored(name)
    offer = evaluate(lines, [])
    header, *body = export_rows(offer, offer_id="offer-1", supplier=result.supplier_name)

    assert header == list(COLUMNS)
    assert len(body) == offer.summary.included_lines

    quantity = header.index("quantity")
    value = header.index("line_value_usd")
    unit_cost = header.index("unit_cost_usd")
    retail_unit = header.index("retail_unit_usd")
    retail_line = header.index("retail_line_usd")

    assert sum(int(r[quantity]) for r in body) == offer.summary.pieces
    assert sum(to_units(Decimal(r[value])) for r in body) == offer.summary.supplier_cost_units
    assert sum(
        to_units(Decimal(r[retail_line])) for r in body if r[retail_line]
    ) == offer.summary.retail_reference_units

    # The retail columns say which is which: unit x pieces == line.
    for row in body:
        assert to_units(Decimal(row[value])) == int(row[quantity]) * to_units(
            Decimal(row[unit_cost])
        )
        if row[retail_unit]:
            assert to_units(Decimal(row[retail_line])) == int(row[quantity]) * to_units(
                Decimal(row[retail_unit])
            )


def test_export_carries_decisions_and_their_notes():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    decisions = save(lines, [], Change("R9", "include", unit_cost="9.00", note="Confirmed by phone"))
    offer = evaluate(lines, decisions)
    header, *body = export_rows(offer, offer_id="o1", supplier="Northstar Supply")

    row = next(r for r in body if r[header.index("source_row")] == "9")
    assert row[header.index("unit_cost_usd")] == "9.00"
    assert "Confirmed by phone" in row[header.index("notes")]


def test_excluded_lines_never_reach_the_file():
    """Northstar gives one line per row, so a row number identifies a line."""
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    header, *body = export_rows(offer, offer_id="o1", supplier="Northstar Supply")

    exported = sorted(int(r[header.index("source_row")]) for r in body)
    assert exported == sorted(line.source_row for line in offer.lines if line.included)
    assert 13 not in exported  # 0 pieces, auto-excluded
    assert 15 not in exported  # unresolved A108 conflict
    assert 9 not in exported   # A104, no supplier cost


@pytest.mark.parametrize("raw", ["=1+1", "+1", "-1", "@SUM(A1)", "\tx", "\rx"])
def test_escape_text_neutralises_every_risky_prefix(raw):
    assert escape_text(raw) == "'" + raw


def test_export_escapes_supplier_controlled_text():
    lines = synthetic(("A100", "M", 10, "3.00"), description='=HYPERLINK("http://evil","click")')
    offer = evaluate(lines, [])
    header, row = export_rows(offer, offer_id="o1", supplier="=cmd|'/c calc'!A1")

    assert row[header.index("description")].startswith("'=HYPERLINK")
    assert row[header.index("supplier")].startswith("'=cmd")
    assert row[header.index("quantity")] == "10"  # numbers stay numbers
    assert row[header.index("unit_cost_usd")] == "3.00"


def test_csv_opens_cleanly_in_excel():
    result, lines = stored("02-harbor-size-grid.xlsx")
    data = to_csv(evaluate(lines, []), offer_id="o1", supplier=result.supplier_name)
    assert data.startswith(b"\xef\xbb\xbf")
    assert data.decode("utf-8-sig").splitlines()[0] == ",".join(COLUMNS)


# ---------- reviewers adding bad rows to the "all valid" file ----------

DAMAGE = [
    ("item_code", None, "MISSING_ITEM_CODE"),
    ("quantity", "TBD", "INVALID_QUANTITY"),
    ("quantity", -5, "NEGATIVE_QUANTITY"),
    ("unit_cost", "$x", "INVALID_COST"),
]


def test_bad_rows_injected_into_the_5000_row_file_are_caught_and_the_rest_stays_exact():
    """CLAUDE.md warns the 5,000-row file claims every row is valid and that
    reviewers will change that. The same rules have to catch them, and the
    untouched rows have to total exactly as before."""
    source = (FIXTURES / "03-northstar-5000-rows.xlsx").read_bytes()
    meta = read_workbook(source)
    clean = evaluate(store(analyze(meta.lines)), [])

    workbook = load_workbook(io.BytesIO(source))
    sheet = workbook[meta.sheet_name]
    rows = [meta.header_row + 1 + i for i in range(len(DAMAGE))]
    for row, (field, value, _) in zip(rows, DAMAGE):
        sheet[f"{meta.columns[field]}{row}"] = value
    buffer = io.BytesIO()
    workbook.save(buffer)

    damaged = read_workbook(buffer.getvalue())
    offer = evaluate(store(analyze(damaged.lines)), [])
    by_line, clean_by_line = offer.by_id(), clean.by_id()

    lost_pieces = lost_cost = lost_retail = 0
    for row, (_, _, code) in zip(rows, DAMAGE):
        line = by_line[f"R{row}"]
        assert not line.included, f"row {row} should have been caught"
        assert code in issue_codes(line)
        was = clean_by_line[f"R{row}"]
        lost_pieces += was.quantity
        lost_cost += was.line_value_units
        lost_retail += was.quantity * was.retail_units

    assert offer.summary.included_lines == 5_000 - len(DAMAGE)
    assert offer.summary.pieces == clean.summary.pieces - lost_pieces
    assert offer.summary.supplier_cost_units == clean.summary.supplier_cost_units - lost_cost
    assert offer.summary.retail_reference_units == clean.summary.retail_reference_units - lost_retail

    # And the export still agrees with the screen after the damage.
    header, *body = export_rows(offer, offer_id="o1", supplier=damaged.supplier_name)
    assert sum(int(r[header.index("quantity")]) for r in body) == offer.summary.pieces
    assert sum(
        to_units(Decimal(r[header.index("line_value_usd")])) for r in body
    ) == offer.summary.supplier_cost_units


# ---------- summary counts drive the tabs ----------

def test_summary_counts_match_the_lines_behind_them():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    summary = offer.summary

    assert summary.total_lines == len(offer.lines)
    assert summary.needs_decision_open == sum(1 for l in offer.lines if l.needs_decision_open)
    assert summary.warnings == sum(1 for l in offer.lines if "warning" in {i.kind for i in l.issues})
    assert summary.fixed == sum(1 for l in offer.lines if "fixed" in {i.kind for i in l.issues})
    assert summary.included_lines == sum(1 for l in offer.lines if l.included)
    assert summary.auto_excluded == 1  # the 0-pieces row
    assert summary.user_excluded == 0


def test_a_decision_closes_the_open_question():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    before = evaluate(lines, []).summary
    after = evaluate(lines, save(lines, [], Change("R12", "exclude"))).summary

    assert after.needs_decision_open == before.needs_decision_open - 1
    assert after.user_excluded == 1
    assert after.pieces == before.pieces  # excluding an open question changes no total


def test_evaluate_returns_the_same_totals_when_run_twice():
    """Nothing in evaluate mutates its input, so a retry can never double."""
    _, lines = stored("02-harbor-size-grid.xlsx")
    decisions = save(lines, [], Change("R13-S", "include", unit_cost="2.00"))
    first: EvaluatedOffer = evaluate(lines, decisions)
    second: EvaluatedOffer = evaluate(lines, decisions)
    assert first.summary == second.summary


# ---------- the workbook ----------

def read_workbook_rows(data: bytes):
    """Header labels plus each row as a dict keyed by column label."""
    book = load_workbook(io.BytesIO(data))
    sheet = book[SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=False))
    headers = [c.value for c in rows[0]]
    body = [{headers[i]: cell for i, cell in enumerate(row)} for row in rows[1:]]
    return headers, body, sheet


@pytest.mark.parametrize("name", ALL_FIXTURES)
def test_workbook_sums_equal_the_summary(name):
    result, lines = stored(name)
    offer = evaluate(lines, [])
    headers, body, _ = read_workbook_rows(
        to_xlsx(offer, offer_id="offer-1", supplier=result.supplier_name)
    )

    assert headers == [label for _, label, _, _ in COLUMN_SPEC]
    assert len(body) == offer.summary.included_lines

    pieces = sum(row["Pieces"].value for row in body)
    cost = sum(to_units(read_amount(row["Line value"].value)) for row in body)
    retail = sum(
        to_units(read_amount(row["Retail line (reference only)"].value))
        for row in body
        if row["Retail line (reference only)"].value is not None
    )
    assert pieces == offer.summary.pieces
    assert cost == offer.summary.supplier_cost_units
    assert retail == offer.summary.retail_reference_units


@pytest.mark.parametrize("name", ALL_FIXTURES)
def test_every_workbook_amount_survives_the_round_trip(name):
    """A spreadsheet number is a double; check that costs us nothing here."""
    result, lines = stored(name)
    offer = evaluate(lines, [])
    records = {
        (r.item_code, r.size, r.source_row): r
        for r in export_records(offer, offer_id="o1", supplier=result.supplier_name)
    }
    _, body, _ = read_workbook_rows(
        to_xlsx(offer, offer_id="o1", supplier=result.supplier_name)
    )
    for row in body:
        key = (row["Item code"].value, row["Size"].value, row["Source row"].value)
        record = records[key]
        assert to_units(read_amount(row["Supplier cost / piece"].value)) == record.unit_cost_units
        assert to_units(read_amount(row["Line value"].value)) == record.line_value_units


def test_leading_zeros_survive_the_workbook():
    """The trap this whole app exists to survive, at the last step: Excel
    reads a CSV column of digits as a number and 000101 opens as 101."""
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    _, body, _ = read_workbook_rows(to_xlsx(offer, offer_id="o1", supplier="Northstar"))

    coded = [row["Item code"] for row in body if row["Item code"].value == "000101"]
    assert coded, "000101 is not in the export at all"
    for cell in coded:
        assert isinstance(cell.value, str)      # not the number 101
        assert cell.data_type == "s"
        assert cell.number_format == "@"


def test_workbook_never_writes_a_live_formula():
    lines = synthetic(("A100", "M", 10, "3.00"), description='=HYPERLINK("http://evil","click")')
    offer = evaluate(lines, [])
    _, body, _ = read_workbook_rows(
        to_xlsx(offer, offer_id="o1", supplier="=cmd|'/c calc'!A1")
    )
    row = body[0]
    for label in ("Description", "Supplier"):
        assert row[label].data_type == "s", f"{label} became a formula"
        # Written literally, so no apostrophe is needed to defuse it.
        assert not str(row[label].value).startswith("'")
    assert row["Description"].value.startswith("=HYPERLINK")
    assert row["Pieces"].value == 10           # numbers stay numbers


def test_workbook_labels_retail_as_reference():
    labels = [label for _, label, _, _ in COLUMN_SPEC]
    retail = [label for label in labels if "Retail" in label]
    assert len(retail) == 2
    assert all("reference only" in label for label in retail)


def test_workbook_sizes_cannot_be_read_as_dates():
    """A size of 1/2 is a size, not the first of February."""
    lines = synthetic(("A100", "1/2", 10, "3.00"))
    offer = evaluate(lines, [])
    _, body, _ = read_workbook_rows(to_xlsx(offer, offer_id="o1", supplier="S"))
    assert body[0]["Size"].value == "1/2"
    assert body[0]["Size"].number_format == "@"


# ---------- the would-be line value ----------

def test_excluded_lines_carry_a_would_be_value_that_no_total_uses():
    """A card comparing conflicting rows needs each row's value, and both rows
    are out until someone picks. The server sends it; nothing sums it."""
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, [])
    by_id = offer.by_id()

    assert by_id["R14"].line_value_units is None
    assert by_id["R14"].would_be_value_units == 100 * to_units(Decimal("5.00"))
    assert by_id["R15"].would_be_value_units == 100 * to_units(Decimal("5.50"))
    # Nothing is invented where a number is missing or unusable.
    assert by_id["R9"].would_be_value_units is None    # no supplier cost
    assert by_id["R12"].would_be_value_units is None   # -12 pieces
    assert by_id["R13"].would_be_value_units is None   # 0 pieces
    assert by_id["R18"].would_be_value_units is None   # "TBD"

    for line in offer.lines:
        if line.included:
            assert line.would_be_value_units == line.line_value_units
    assert offer.summary.supplier_cost_units == to_units(Decimal("4208.00"))


# ---------- notes worded for a file ----------

SCREEN_WORDS = ("Enter one", "Choose which", "Include only if", "Check the original",
                "Counted once here")


def _decided(lines):
    """Every kind of decision at once: pick, count once, fix, exclude."""
    return save(
        lines, [],
        Change("R14", "include"), Change("R15", "exclude"),
        Change("R6", "include"), Change("R10", "exclude"),
        Change("R9", "include", unit_cost="9.00", note="Confirmed by phone"),
        Change("R12", "exclude", note="Supplier to confirm"),
        Change("R16", "include", item_code="SC-01"),
        Change("R18", "exclude"),
    )


@pytest.mark.parametrize("name", ALL_FIXTURES[:2])
def test_file_notes_carry_no_screen_instructions(name):
    _, lines = stored(name)
    decisions = _decided(lines) if name.startswith("01") else []
    for offer in (evaluate(lines, []), evaluate(lines, decisions)):
        texts = [r.notes for r in export_records(offer, offer_id="o1", supplier="S")]
        texts += [r.reason for r in left_out_records(offer)]
        for text in texts:
            assert not any(word in text for word in SCREEN_WORDS), text
            assert ".;" not in text and not text.endswith("."), text


def test_the_stripped_instructions_really_occur_in_issue_messages():
    """Otherwise the test above would pass by never meeting one."""
    _, lines = stored("01-northstar-line-sheet.xlsx")
    messages = " ".join(i.message for line in evaluate(lines, []).lines for i in line.issues)
    for word in SCREEN_WORDS:
        if word != "Check the original":   # LARGE_QUANTITY; no fixture has one
            assert word in messages, word


def test_file_notes_follow_the_saved_decisions():
    _, lines = stored("01-northstar-line-sheet.xlsx")

    def notes(*changes):
        offer = evaluate(lines, save(lines, [], *changes))
        return {r.source_row: r.notes for r in export_records(offer, offer_id="o1", supplier="S")}

    default = notes()
    assert default[7] == "Supplier cost: sheet had '$3.25', read as 3.25"
    assert default[19] == "Pieces: sheet had ' 45 ', read as 45"
    assert default[6] == "Identical to row 10, which is left out so the stock is counted once"

    # Counting both makes "counted once" untrue, so it must not be said.
    both = notes(Change("R6", "include"), Change("R10", "include"))
    assert "counted once" not in both[6]
    assert "also included on row(s) 10" in both[6]
    assert both[10].endswith("Included by reviewer as a separate lot from row 6")

    chosen = notes(Change("R14", "include"), Change("R15", "exclude"))
    assert chosen[14] == "Chosen by reviewer over row 15, which had different values"

    # A typed value says it was typed, and what the sheet held.
    typed = notes(Change("R9", "include", unit_cost="9.00", note="Confirmed by phone"))
    assert typed[9] == (
        "Included by reviewer; Supplier cost set by reviewer to 9.00 (sheet: blank); "
        "Reviewer's note: Confirmed by phone"
    )


# ---------- the Summary sheet ----------

def read_summary(data: bytes):
    """The Summary sheet as {label: cell} plus the left-out table rows."""
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == [SUMMARY_SHEET, SHEET_NAME]
    rows = list(book[SUMMARY_SHEET].iter_rows())
    labels = {row[0].value: row[1] for row in rows if isinstance(row[0].value, str)}
    start = next(
        (i for i, row in enumerate(rows) if row[0].value == "Why it isn't included"), None
    )
    table = [] if start is None else [
        [c.value for c in row] for row in rows[start + 1:] if row[0].value is not None
    ]
    return labels, table


def test_summary_sheet_says_what_the_file_is():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    offer = evaluate(lines, save(lines, [], Change("R12", "exclude", note="Supplier to confirm")))
    url = "https://offers.example/offers/offer-1"
    labels, table = read_summary(to_xlsx(
        offer, offer_id="offer-1", supplier="Northstar Supply",
        context=WorkbookContext(
            offer_url=url, source_filename="01-northstar-line-sheet.xlsx",
            sheet_name="Line Sheet", version=2,
            exported_at=datetime(2026, 9, 22, 14, 3, tzinfo=timezone.utc),
        ),
    ))
    summary = offer.summary

    assert labels["Offer link"].value == url
    assert labels["Offer link"].hyperlink.target == url
    assert labels["Offer ID"].value == "offer-1"
    assert labels["Supplier"].value == "Northstar Supply"
    assert labels["Source file"].value == "01-northstar-line-sheet.xlsx"
    assert labels["Saved version"].value == 2
    assert labels["Exported at"].value == "2026-09-22 14:03 UTC"

    assert labels["Lines included"].value == summary.included_lines
    assert labels["Pieces"].value == summary.pieces
    assert labels["Pieces"].number_format == "#,##0"
    assert to_units(read_amount(labels["Supplier cost total"].value)) == summary.supplier_cost_units
    assert to_units(read_amount(
        labels["Retail reference total (not what we pay)"].value
    )) == summary.retail_reference_units
    awaiting = sum(1 for line in offer.lines if line.needs_decision_open)
    assert labels["Lines left out"].value == (
        summary.total_lines - summary.included_lines - awaiting
    )
    assert labels["Awaiting a decision"].value == awaiting
    # The screen splits these two the same way, so the numbers can be compared
    # without reconciling anything.
    assert labels["Lines left out"].value + awaiting == len(table)

    # Every line out of the totals is listed, once, with its reason.
    reasons = {row[1]: row[0] for row in table}
    assert sorted(reasons) == sorted(l.source_row for l in offer.lines if not l.included)
    assert reasons[12] == "Excluded by reviewer: Supplier to confirm"
    assert reasons[9] == "Awaiting a decision: No supplier cost"
    assert reasons[13] == "0 pieces available"


@pytest.mark.parametrize("name", ALL_FIXTURES)
def test_summary_sheet_totals_equal_the_summary(name):
    result, lines = stored(name)
    offer = evaluate(lines, [])
    labels, table = read_summary(to_xlsx(offer, offer_id="o1", supplier=result.supplier_name))
    summary = offer.summary

    assert labels["Pieces"].value == summary.pieces
    assert to_units(read_amount(labels["Supplier cost total"].value)) == summary.supplier_cost_units
    awaiting = sum(1 for line in offer.lines if line.needs_decision_open)
    assert labels["Lines left out"].value == len(table) - awaiting
    assert labels["Awaiting a decision"].value == awaiting
    assert len(table) == summary.total_lines - summary.included_lines
    if not table:
        assert "None: every line in the sheet is included." in labels


def test_offer_id_lives_on_the_summary_not_on_every_row():
    assert "Offer ID" not in [label for _, label, _, _ in COLUMN_SPEC]


def test_pieces_are_formatted_as_counts():
    _, lines = stored("01-northstar-line-sheet.xlsx")
    _, body, _ = read_workbook_rows(to_xlsx(evaluate(lines, []), offer_id="o1", supplier="S"))
    assert body and all(row["Pieces"].number_format == "#,##0" for row in body)
    assert all(row["Source row"].number_format == "General" for row in body)
