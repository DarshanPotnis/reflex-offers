"""Storage and API behaviour.

Two things are being proved here: the answer key survives a round trip
through the database, and saved work cannot be lost or doubled.
"""

import io
import re
import uuid
from contextlib import contextmanager
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app import service
from app.core.export_xlsx import read_amount
from app.core.money import to_units
from app.models import DecisionRow, Offer, SaveRequest
from tests.conftest import fixture_bytes, run_together, upload

ANSWER_KEY = [
    ("01-northstar-line-sheet.xlsx", "northstar", 6, 1_733, "4208.00", "21552.00"),
    ("02-harbor-size-grid.xlsx", "harbor", 22, 1_340, "3740.00", "16144.00"),
    ("03-northstar-5000-rows.xlsx", "northstar", 5_000, 62_444, "187214.50", "1031508.00"),
]


def new_request_id() -> str:
    return str(uuid.uuid4())


def save(client, offer_id, changes, *, base_version=1, request_id=None, headers=None):
    return client.post(
        f"/api/offers/{offer_id}/decisions",
        json={
            "request_id": request_id or new_request_id(),
            "base_version": base_version,
            "changes": changes,
        },
        headers=headers or {},
    )


def count(sessions, model, offer_id) -> int:
    with sessions() as session:
        return len([
            row for row in session.query(model).filter_by(offer_id=offer_id).all()
        ])


# ---------- upload and read ----------

@pytest.mark.parametrize("name, layout, lines, pieces, cost, retail", ANSWER_KEY)
def test_uploaded_offer_matches_the_answer_key(client, name, layout, lines, pieces, cost, retail):
    created = upload(client, name)
    assert created.status_code == 201, created.text
    offer_id = created.json()["id"]

    body = client.get(f"/api/offers/{offer_id}").json()
    assert body["layout"] == layout
    assert body["version"] == 1
    summary = body["summary"]
    assert summary["included_lines"] == lines
    assert summary["pieces"] == pieces
    assert summary["supplier_cost"] == cost
    assert summary["retail_reference"] == retail


def test_money_is_never_a_float_in_json(client):
    upload(client, "01-northstar-line-sheet.xlsx")
    offer_id = client.get("/api/offers").json()["offers"][0]["id"]
    body = client.get(f"/api/offers/{offer_id}").json()

    assert isinstance(body["summary"]["supplier_cost"], str)
    for line in body["lines"]:
        for field in ("unit_cost", "retail", "line_value"):
            assert line[field] is None or isinstance(line[field], str)


def test_lines_carry_the_original_cells_for_checking_against_the_sheet(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    lines = {ln["line_id"]: ln for ln in client.get(f"/api/offers/{offer_id}").json()["lines"]}

    cost_cell = lines["R7"]["cells"]["unit_cost"]
    assert cost_cell["raw"] == "$3.25"          # what the supplier typed
    assert re.fullmatch(r"[A-Z]+7", cost_cell["coordinate"])  # where to look
    assert lines["R7"]["unit_cost"] == "3.25"  # what we will pay
    assert lines["R7"]["original"]["unit_cost_units"]["status"] == "normalized"


def test_notices_survive_the_round_trip(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    notices = client.get(f"/api/offers/{offer_id}").json()["notices"]
    assert any("Row 21 skipped" in n for n in notices)


def test_unknown_offer_is_404(client):
    assert client.get(f"/api/offers/{uuid.uuid4()}").status_code == 404
    assert client.get(f"/api/offers/{uuid.uuid4()}/export.csv").status_code == 404
    assert save(client, str(uuid.uuid4()), [{"line_id": "R1", "action": "exclude"}]).status_code == 404


def test_unsupported_file_is_refused_with_a_clear_message(client):
    response = client.post(
        "/api/offers", files={"file": ("notes.txt", b"not a spreadsheet", "text/plain")}
    )
    assert response.status_code == 422
    assert "Excel workbook" in response.json()["detail"]


def test_upload_over_the_limit_is_refused(client):
    oversized = b"x" * (service.MAX_UPLOAD_BYTES + 1)
    response = client.post("/api/offers", files={"file": ("big.xlsx", oversized)})
    assert response.status_code == 413


def test_original_file_is_kept_byte_for_byte(client):
    name = "02-harbor-size-grid.xlsx"
    offer_id = upload(client, name).json()["id"]
    assert client.get(f"/api/offers/{offer_id}/source").content == fixture_bytes(name)


def test_server_timing_locates_the_slowest_stage(client):
    created = upload(client, "01-northstar-line-sheet.xlsx")
    assert "parse" in created.headers["server-timing"]
    offer_id = created.json()["id"]
    timing = client.get(f"/api/offers/{offer_id}").headers["server-timing"]
    assert "db" in timing and "evaluate" in timing and "serialize" in timing


# ---------- upload idempotency ----------

def test_same_idempotency_key_makes_one_offer(client):
    first = upload(client, "01-northstar-line-sheet.xlsx", key="abc-123")
    second = upload(client, "01-northstar-line-sheet.xlsx", key="abc-123")

    assert first.status_code == 201 and second.status_code == 200
    assert first.json()["id"] == second.json()["id"]
    assert len(client.get("/api/offers").json()["offers"]) == 1


def test_same_idempotency_key_with_a_different_file_is_refused(client):
    upload(client, "01-northstar-line-sheet.xlsx", key="abc-123")
    clash = upload(client, "02-harbor-size-grid.xlsx", key="abc-123")

    assert clash.status_code == 409
    assert len(client.get("/api/offers").json()["offers"]) == 1


def test_two_uploads_without_a_key_stay_separate(client):
    first = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    second = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    assert first != second

    save(client, first, [{"line_id": "R13", "action": "include", "quantity": 5}])

    assert client.get(f"/api/offers/{first}").json()["version"] == 2
    assert client.get(f"/api/offers/{second}").json()["version"] == 1
    assert client.get(f"/api/offers/{second}").json()["summary"]["pieces"] == 1_733


def test_recent_offers_are_capped(client, monkeypatch):
    monkeypatch.setattr(service, "RECENT_OFFER_LIMIT", 3)
    for _ in range(5):
        upload(client, "01-northstar-line-sheet.xlsx")
    assert len(client.get("/api/offers").json()["offers"]) == 3


# ---------- saving decisions ----------

def test_saving_a_decision_bumps_the_version_and_changes_the_totals(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]

    receipt = save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}])
    assert receipt.status_code == 200
    assert receipt.json()["version"] == 2
    assert receipt.json()["applied"] == 1

    body = client.get(f"/api/offers/{offer_id}").json()
    assert body["version"] == 2
    assert body["summary"]["pieces"] == 1_833
    assert body["summary"]["supplier_cost"] == "4708.00"


def test_keep_this_one_resolves_a_whole_group_in_one_save(client):
    """CLAUDE.md: resolving a conflict must not leave a sibling open."""
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    before = client.get(f"/api/offers/{offer_id}").json()["summary"]["needs_decision_open"]

    response = save(client, offer_id, [
        {"line_id": "R14", "action": "include", "unit_cost": "5.00"},
        {"line_id": "R15", "action": "exclude"},
    ])
    assert response.status_code == 200
    assert response.json()["version"] == 2  # one bump for the whole group

    body = client.get(f"/api/offers/{offer_id}").json()
    lines = {ln["line_id"]: ln for ln in body["lines"]}
    assert lines["R14"]["status"] == "included"
    assert lines["R15"]["status"] == "excluded"
    assert not lines["R15"]["needs_decision_open"]
    assert body["summary"]["needs_decision_open"] == before - 2


def test_decisions_survive_a_new_connection(client, sessions):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    save(client, offer_id, [{"line_id": "R9", "action": "include", "unit_cost": "9.00"}])

    with sessions() as session:
        _, evaluated = service.load_evaluated(session, offer_id)
    assert evaluated.by_id()["R9"].included
    assert evaluated.summary.pieces == 1_793


def test_reset_removes_the_decision_row(client, sessions):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}])
    assert count(sessions, DecisionRow, offer_id) == 1

    save(client, offer_id, [{"line_id": "R14", "action": "reset"}], base_version=2)
    assert count(sessions, DecisionRow, offer_id) == 0
    assert client.get(f"/api/offers/{offer_id}").json()["summary"]["pieces"] == 1_733


# ---------- retries and versions ----------

def test_replaying_a_request_id_changes_nothing(client, sessions):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    request_id = new_request_id()
    changes = [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}]

    first = save(client, offer_id, changes, request_id=request_id)
    second = save(client, offer_id, changes, request_id=request_id)

    assert first.json() == second.json()
    assert second.headers["x-replayed"] == "true"
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2
    assert count(sessions, DecisionRow, offer_id) == 1
    assert client.get(f"/api/offers/{offer_id}").json()["summary"]["pieces"] == 1_833


def test_request_id_reused_with_different_changes_is_refused(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    request_id = new_request_id()

    save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}],
         request_id=request_id)
    clash = save(client, offer_id, [{"line_id": "R15", "action": "exclude"}],
                 request_id=request_id)

    assert clash.status_code == 422
    assert "reused" in clash.json()["errors"][0]["messages"][0]
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2


def test_stale_base_version_is_refused_and_saves_nothing(client, sessions):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}])

    stale = save(client, offer_id, [{"line_id": "R12", "action": "exclude"}], base_version=1)

    assert stale.status_code == 409
    assert stale.json()["current_version"] == 2
    assert count(sessions, DecisionRow, offer_id) == 1
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2


def test_two_saves_on_the_same_base_version_only_one_wins(client, sessions):
    """The deterministic half of the race: whatever the timing, the second
    save must not silently overwrite the first."""
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]

    first = save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}],
                 base_version=1)
    second = save(client, offer_id, [{"line_id": "R12", "action": "exclude"}], base_version=1)

    assert sorted([first.status_code, second.status_code]) == [200, 409]
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2
    assert count(sessions, DecisionRow, offer_id) == 1


# ---------- concurrency (real threads) ----------

def _save_job(offer_id, request_id, changes, base_version=1):
    from app.api import ChangeIn, DecisionsIn, canonical_body_hash

    payload = DecisionsIn(
        request_id=request_id, base_version=base_version,
        changes=[ChangeIn(**c) for c in changes],
    )

    def job(session):
        return service.save_decisions(
            session,
            offer_id=offer_id,
            request_id=str(payload.request_id),
            base_version=payload.base_version,
            changes=[c.to_core() for c in payload.changes],
            body_sha256=canonical_body_hash(payload),
        )

    return job


@pytest.mark.parametrize("_run", [1, 2, 3])
def test_concurrent_saves_on_one_base_version(client, sessions, is_postgres, _run):
    if not is_postgres:
        pytest.skip("SQLite serialises writers; this race only exists on Postgres")

    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    results = run_together(sessions, [
        _save_job(offer_id, new_request_id(), [{"line_id": "R14", "action": "include",
                                                "unit_cost": "5.00"}]),
        _save_job(offer_id, new_request_id(), [{"line_id": "R12", "action": "exclude"}]),
    ])

    kinds = sorted(kind for kind, _ in results)
    assert kinds == ["error", "ok"], results
    conflict = next(value for kind, value in results if kind == "error")
    assert isinstance(conflict, service.VersionConflict)
    assert conflict.current_version == 2

    with sessions() as session:
        assert session.get(Offer, offer_id).version == 2
    assert count(sessions, DecisionRow, offer_id) == 1


@pytest.mark.parametrize("_run", [1, 2, 3])
def test_concurrent_duplicate_retries_apply_once(client, sessions, is_postgres, _run):
    """Two copies of one retry. Both must get the same receipt, and the work
    must happen exactly once."""
    if not is_postgres:
        pytest.skip("SQLite serialises writers; this race only exists on Postgres")

    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    request_id = new_request_id()
    changes = [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}]

    results = run_together(sessions, [
        _save_job(offer_id, request_id, changes),
        _save_job(offer_id, request_id, changes),
    ])

    assert all(kind == "ok" for kind, _ in results), results
    receipts = [value[0] for _, value in results]
    assert receipts[0] == receipts[1]
    assert receipts[0]["version"] == 2

    with sessions() as session:
        assert session.get(Offer, offer_id).version == 2
    assert count(sessions, DecisionRow, offer_id) == 1
    assert count(sessions, SaveRequest, offer_id) == 1
    assert client.get(f"/api/offers/{offer_id}").json()["summary"]["pieces"] == 1_833


def test_a_duplicate_retry_that_loses_the_version_race_still_gets_its_receipt(
    client, sessions, monkeypatch
):
    """The same recovery as above, forced deterministically so it runs on any
    database: the twin's receipt lands between this save's first lookup and
    its conditional update, so the update finds nothing to bump.

    Without the re-check at that point, this save would return 409 for work
    its own twin had just committed.
    """
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    request_id = new_request_id()
    changes = [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}]

    first = save(client, offer_id, changes, request_id=request_id)
    assert first.status_code == 200

    real_replay = service._replay
    calls = {"n": 0}

    def blind_first_lookup(*args, **kwargs):
        calls["n"] += 1
        return None if calls["n"] == 1 else real_replay(*args, **kwargs)

    monkeypatch.setattr(service, "_replay", blind_first_lookup)

    with sessions() as session:
        receipt, replayed = _save_job(offer_id, request_id, changes)(session)

    assert calls["n"] == 2, "the step-3 re-check never ran"
    assert replayed and receipt == first.json()
    assert count(sessions, DecisionRow, offer_id) == 1
    assert count(sessions, SaveRequest, offer_id) == 1
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2


# ---------- fault injection ----------

@pytest.fixture
def faults(monkeypatch):
    monkeypatch.setenv("FAULT_INJECTION", "1")


def test_after_commit_failure_then_retry_applies_exactly_once(client, sessions, faults):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    request_id = new_request_id()
    changes = [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}]

    failed = save(client, offer_id, changes, request_id=request_id,
                  headers={"X-Fault": "after-commit"})
    assert failed.status_code == 503

    # The dangerous state: saved, but the client was told it failed.
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 2

    retry = save(client, offer_id, changes, request_id=request_id)
    assert retry.status_code == 200
    assert retry.json()["version"] == 2
    assert retry.headers["x-replayed"] == "true"

    body = client.get(f"/api/offers/{offer_id}").json()
    assert body["version"] == 2  # +1 in total, not +2
    assert body["summary"]["pieces"] == 1_833
    assert count(sessions, DecisionRow, offer_id) == 1


def test_before_commit_failure_saves_nothing(client, sessions, faults):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    changes = [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}]

    failed = save(client, offer_id, changes, headers={"X-Fault": "before-commit"})
    assert failed.status_code == 503

    body = client.get(f"/api/offers/{offer_id}").json()
    assert body["version"] == 1
    assert body["summary"]["pieces"] == 1_733
    assert count(sessions, DecisionRow, offer_id) == 0
    assert count(sessions, SaveRequest, offer_id) == 0


def test_fault_headers_are_ignored_unless_enabled(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    response = save(client, offer_id, [{"line_id": "R14", "action": "include",
                                        "unit_cost": "5.00"}],
                    headers={"X-Fault": "after-commit"})
    assert response.status_code == 200


# ---------- data protection: invalid values sent straight to the API ----------

BAD_CHANGES = [
    ({"line_id": "R9", "action": "include", "quantity": -5, "unit_cost": "9.00"}, "between 1 and"),
    ({"line_id": "R9", "action": "include", "quantity": 0, "unit_cost": "9.00"}, "between 1 and"),
    ({"line_id": "R9", "action": "include", "quantity": 12.5, "unit_cost": "9.00"}, "whole number"),
    ({"line_id": "R9", "action": "include", "quantity": "12", "unit_cost": "9.00"}, "whole number"),
    ({"line_id": "R9", "action": "include", "quantity": True, "unit_cost": "9.00"}, "whole number"),
    ({"line_id": "R9", "action": "include", "unit_cost": "abc"}, "isn't an amount"),
    ({"line_id": "R9", "action": "include", "unit_cost": "-1"}, "above $0"),
    ({"line_id": "R9", "action": "include", "unit_cost": "0"}, "above $0"),
    ({"line_id": "R9", "action": "include", "unit_cost": "1.23456"}, "isn't an amount"),
    ({"line_id": "R9", "action": "include", "unit_cost": 9.0}, "decimal string"),
    ({"line_id": "R9", "action": "include", "unit_cost": "2000000"}, "limit"),
    ({"line_id": "R9", "action": "include"}, "no usable supplier cost"),
    ({"line_id": "R16", "action": "include", "unit_cost": "1.00"}, "no item code"),
    ({"line_id": "nope", "action": "include"}, "no line"),
    ({"line_id": "R9", "action": "exclude", "quantity": 5}, "can't be set while excluding"),
    ({"line_id": "R9", "action": "sideways"}, "action"),
    ({"line_id": "R9", "action": "include", "colour": "red"}, "Unknown field"),
]


@pytest.mark.parametrize("change, expected", BAD_CHANGES)
def test_invalid_values_are_refused_and_the_offer_is_unchanged(client, change, expected):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    before = client.get(f"/api/offers/{offer_id}").json()

    response = save(client, offer_id, [change])

    assert response.status_code == 422, response.text
    messages = " ".join(
        m for e in response.json()["errors"] for m in e["messages"]
    )
    assert expected.lower() in messages.lower(), messages
    assert client.get(f"/api/offers/{offer_id}").json() == before


def test_errors_are_reported_against_their_line(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    response = save(client, offer_id, [
        {"line_id": "R9", "action": "include", "quantity": 12.5, "unit_cost": "9.00"},
    ])
    assert response.status_code == 422
    assert response.json()["errors"][0]["line_id"] == "R9"


def test_one_bad_change_in_a_batch_applies_none(client, sessions):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    before = client.get(f"/api/offers/{offer_id}").json()

    response = save(client, offer_id, [
        {"line_id": "R14", "action": "include", "unit_cost": "5.00"},
        {"line_id": "R9", "action": "include", "unit_cost": "abc"},
        {"line_id": "R12", "action": "exclude"},
    ])

    assert response.status_code == 422
    assert count(sessions, DecisionRow, offer_id) == 0
    assert client.get(f"/api/offers/{offer_id}").json() == before


def test_empty_changes_is_refused(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    response = save(client, offer_id, [])
    assert response.status_code == 422
    assert "no changes" in response.json()["errors"][0]["messages"][0]
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 1


def test_the_same_line_twice_in_one_save_is_refused(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    response = save(client, offer_id, [
        {"line_id": "R14", "action": "include", "unit_cost": "5.00"},
        {"line_id": "R14", "action": "exclude"},
    ])
    assert response.status_code == 422
    assert "twice" in response.json()["errors"][0]["messages"][0]
    assert client.get(f"/api/offers/{offer_id}").json()["version"] == 1


# ---------- export ----------

@pytest.mark.parametrize("name, layout, lines, pieces, cost, retail", ANSWER_KEY)
def test_export_agrees_with_the_screen(client, name, layout, lines, pieces, cost, retail):
    offer_id = upload(client, name).json()["id"]
    summary = client.get(f"/api/offers/{offer_id}").json()["summary"]

    response = client.get(f"/api/offers/{offer_id}/export.csv")
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")

    rows = response.content.decode("utf-8-sig").strip().splitlines()
    header = rows[0].split(",")
    body = [r.split(",") for r in rows[1:]]

    assert len(body) == summary["included_lines"] == lines
    assert sum(int(r[header.index("quantity")]) for r in body) == summary["pieces"] == pieces
    assert sum(
        to_units(Decimal(r[header.index("line_value_usd")])) for r in body
    ) == to_units(Decimal(summary["supplier_cost"]))


def test_export_reflects_saved_decisions(client):
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    save(client, offer_id, [{"line_id": "R14", "action": "include", "unit_cost": "5.00"}])

    text = client.get(f"/api/offers/{offer_id}/export.csv").content.decode("utf-8-sig")
    summary = client.get(f"/api/offers/{offer_id}").json()["summary"]
    rows = text.strip().splitlines()
    header = rows[0].split(",")

    assert sum(
        int(r.split(",")[header.index("quantity")]) for r in rows[1:]
    ) == summary["pieces"] == 1_833


# ---------- the workbook endpoint ----------

@pytest.mark.parametrize("name, layout, lines, pieces, cost, retail", ANSWER_KEY)
def test_workbook_export_agrees_with_the_screen(client, name, layout, lines, pieces, cost, retail):
    offer_id = upload(client, name).json()["id"]
    summary = client.get(f"/api/offers/{offer_id}").json()["summary"]

    response = client.get(f"/api/offers/{offer_id}/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert response.content[:2] == b"PK"  # a real zip, i.e. a real xlsx
    assert response.headers["content-disposition"].endswith('.xlsx"')

    book = load_workbook(io.BytesIO(response.content))
    sheet = book["Offer"]
    rows = list(sheet.iter_rows(values_only=True))
    header = list(rows[0])
    body = rows[1:]

    assert len(body) == summary["included_lines"] == lines
    quantity = header.index("Pieces")
    value = header.index("Line value")
    assert sum(r[quantity] for r in body) == summary["pieces"] == pieces
    assert sum(
        to_units(read_amount(r[value])) for r in body
    ) == to_units(Decimal(summary["supplier_cost"]))


def test_workbook_keeps_item_codes_as_text(client):
    """Downloaded and reopened, 000101 is still 000101 and not 101."""
    offer_id = upload(client, "01-northstar-line-sheet.xlsx").json()["id"]
    data = client.get(f"/api/offers/{offer_id}/export.xlsx").content

    sheet = load_workbook(io.BytesIO(data))["Offer"]
    header = [c.value for c in next(sheet.iter_rows())]
    column = header.index("Item code") + 1
    codes = [sheet.cell(row=r, column=column) for r in range(2, sheet.max_row + 1)]

    assert any(cell.value == "000101" for cell in codes)
    assert all(isinstance(cell.value, str) for cell in codes)


def test_both_export_formats_carry_the_same_numbers(client):
    offer_id = upload(client, "02-harbor-size-grid.xlsx").json()["id"]

    csv_rows = client.get(f"/api/offers/{offer_id}/export.csv").content.decode(
        "utf-8-sig"
    ).strip().split("\r\n")
    csv_header = csv_rows[0].split(",")
    csv_pieces = sum(int(r.split(",")[csv_header.index("quantity")]) for r in csv_rows[1:])

    sheet = load_workbook(
        io.BytesIO(client.get(f"/api/offers/{offer_id}/export.xlsx").content)
    )["Offer"]
    rows = list(sheet.iter_rows(values_only=True))
    xlsx_pieces = sum(r[list(rows[0]).index("Pieces")] for r in rows[1:])

    assert csv_pieces == xlsx_pieces
    assert len(rows) - 1 == len(csv_rows) - 1


def test_workbook_for_an_unknown_offer_is_404(client):
    assert client.get(f"/api/offers/{uuid.uuid4()}/export.xlsx").status_code == 404


# ---------- how the 5,000 lines actually reach the database ----------

def test_five_thousand_lines_are_inserted_in_batches(client, db):
    """One executemany, not one round trip per line.

    On SQLite the difference is milliseconds. Over a network it is the
    difference between one round trip and five thousand, so this is asserted
    directly rather than inferred from a stopwatch.
    """
    from sqlalchemy import event

    executions: list[tuple[str, bool, int]] = []

    def record(conn, cursor, statement, parameters, context, executemany):
        verb = statement.strip().split(None, 1)[0].upper()
        rows = len(parameters) if executemany and parameters else 1
        executions.append((verb, executemany, rows))

    event.listen(db, "before_cursor_execute", record)
    try:
        response = upload(client, "03-northstar-5000-rows.xlsx")
    finally:
        event.remove(db, "before_cursor_execute", record)

    assert response.status_code == 201
    offer_id = response.json()["id"]
    assert client.get(f"/api/offers/{offer_id}").json()["summary"]["included_lines"] == 5_000

    inserts = [e for e in executions if e[0] == "INSERT"]
    rows_inserted = sum(rows for _, _, rows in inserts)
    assert rows_inserted >= 5_000, f"only {rows_inserted} rows inserted"
    # The whole upload is a handful of statements, not thousands.
    assert len(inserts) < 25, (
        f"{len(inserts)} INSERT round trips for 5,000 lines — "
        "the bulk insert has regressed to one statement per row"
    )
    print(
        f"\n  5,000 lines -> {len(inserts)} INSERT statement(s), "
        f"{rows_inserted} rows, {len(executions)} statements in total"
    )


# ---------- Server-Timing must not blame CPU for network ----------

def test_no_sql_runs_inside_the_evaluate_stage(client, db, monkeypatch):
    """`evaluate` is pure CPU, and the header has to keep saying so.

    This regressed once: the upload path called load_evaluated() without
    passing its Timings, so re-reading all 5,000 stored lines landed inside
    the caller's `evaluate` stage. On SQLite that hid in the noise; over a
    cross-country Postgres link it showed up as 2.7 seconds of CPU that was
    really network.
    """
    from sqlalchemy import event

    statements: list[str] = []

    def record(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement.strip().split(None, 1)[0].upper())

    sql_during: dict[str, int] = {}

    class WatchedTimings(service.Timings):
        @contextmanager
        def stage(self, name):
            before = len(statements)
            with super().stage(name):
                yield
            sql_during[name] = sql_during.get(name, 0) + (len(statements) - before)

    monkeypatch.setattr(service, "Timings", WatchedTimings)
    event.listen(db, "before_cursor_execute", record)
    try:
        offer_id = upload(client, "03-northstar-5000-rows.xlsx").json()["id"]
        client.get(f"/api/offers/{offer_id}")
        client.get(f"/api/offers/{offer_id}/export.xlsx")
        client.get(f"/api/offers/{offer_id}/export.csv")
    finally:
        event.remove(db, "before_cursor_execute", record)

    assert sql_during.get("evaluate", 0) == 0, (
        f"{sql_during['evaluate']} SQL statement(s) ran inside the evaluate "
        f"stage; it is supposed to be pure CPU. Stages: {sql_during}"
    )
    assert sql_during.get("parse", 0) == 0
    assert sql_during.get("serialize", 0) == 0
    # And the reads really are being counted somewhere.
    assert sql_during.get("db", 0) >= 4, sql_during
