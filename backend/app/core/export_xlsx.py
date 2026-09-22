"""The downloadable workbook — the format an operations person actually opens.

Why this exists as well as the CSV: Excel reads a CSV column of digits as a
number, so item code ``000101`` opens as ``101``. That is precisely the trap
`parse_item_code` works to survive, undone at the last step. Here the code is
written as a **text** cell and stays ``000101``.

Three things this file gets deliberately right:

**Text cells stay text.** openpyxl turns a string starting with ``=`` into a
formula, so a supplier description of ``=HYPERLINK(...)`` would become live on
open. Every text cell is forced back to ``data_type = "s"`` after assignment,
which writes it as an inline string — no apostrophe needed, unlike the CSV.
Item codes and sizes also get the ``@`` (text) number format, so ``000101``
keeps its zeros and a size of ``1/2`` is not read as a date.

**Money is numeric, so the columns can be summed.** A spreadsheet's numeric
cell is an IEEE double — that is the file format, not a choice we get to
make. Our code never builds a float: it hands openpyxl the exact ``Decimal``
from `money.to_decimal` and lets the library serialise it. For everything in
our range this round-trips exactly, and a test reads the workbook back and
checks every amount against the summary. The CSV remains the exact-decimal
artifact for anything that needs to be byte-exact.

**Retail is labelled.** Two columns, both saying "reference only", so a retail
figure can never be mistaken for what we pay.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .evaluate import EvaluatedOffer
from .export import ExportRecord, export_records
from .money import to_decimal

SHEET_NAME = "Offer"

# (field, header, kind, width)
COLUMN_SPEC: tuple[tuple[str, str, str, int], ...] = (
    ("offer_id", "Offer ID", "text", 38),
    ("supplier", "Supplier", "text", 20),
    ("item_code", "Item code", "code", 14),
    ("description", "Description", "text", 30),
    ("size", "Size", "code", 10),
    ("category", "Category", "text", 14),
    ("quantity", "Pieces", "int", 10),
    ("unit_cost_usd", "Supplier cost / piece", "money", 18),
    ("line_value_usd", "Line value", "money", 14),
    ("retail_unit_usd", "Retail / piece (reference only)", "money", 26),
    ("retail_line_usd", "Retail line (reference only)", "money", 26),
    ("source_row", "Source row", "int", 11),
    ("notes", "Notes", "text", 60),
)

# At least two decimals, up to four, so a $0.0125 cost is not shown as $0.01.
MONEY_FORMAT = '"$"#,##0.00##'


def _write_text(sheet: Worksheet, row: int, column: int, value: str, *, as_code: bool) -> None:
    if not value:
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    # Assigning re-infers the type, so force it back: a description of
    # '=cmd|...' must be a string, never a formula.
    cell.data_type = "s"
    if as_code:
        cell.number_format = "@"


def _write_money(sheet: Worksheet, row: int, column: int, units: int | None) -> None:
    if units is None:
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = to_decimal(units)  # exact Decimal in, never a float we made
    cell.number_format = MONEY_FORMAT


def _values(record: ExportRecord) -> dict[str, object]:
    return {
        "offer_id": record.offer_id,
        "supplier": record.supplier,
        "item_code": record.item_code,
        "description": record.description,
        "size": record.size,
        "category": record.category,
        "quantity": record.quantity,
        "unit_cost_usd": record.unit_cost_units,
        "line_value_usd": record.line_value_units,
        "retail_unit_usd": record.retail_unit_units,
        "retail_line_usd": record.retail_line_units,
        "source_row": record.source_row,
        "notes": record.notes,
    }


def to_xlsx(offer: EvaluatedOffer, *, offer_id: str, supplier: str | None) -> bytes:
    """Included lines only, from saved state. No totals row."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME

    for index, (_, header, _, width) in enumerate(COLUMN_SPEC, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.value = header
        cell.data_type = "s"
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"

    for offset, record in enumerate(
        export_records(offer, offer_id=offer_id, supplier=supplier), start=2
    ):
        values = _values(record)
        for index, (field, _, kind, _width) in enumerate(COLUMN_SPEC, start=1):
            value = values[field]
            if kind == "money":
                _write_money(sheet, offset, index, value)  # type: ignore[arg-type]
            elif kind == "int":
                sheet.cell(row=offset, column=index).value = value
            else:
                _write_text(
                    sheet, offset, index, str(value), as_code=(kind == "code")
                )

    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(COLUMN_SPEC))}{max(sheet.max_row, 1)}"
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def read_amount(value: object) -> Decimal | None:
    """Read a money cell back as an exact Decimal.

    openpyxl hands back a float, because that is what the cell holds. The
    shortest round-trip representation of that double is the decimal we
    wrote; going via ``repr`` recovers it without inventing digits.
    """
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    return Decimal(repr(value))
