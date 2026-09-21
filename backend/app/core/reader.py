"""Read a supplier workbook into standard lines.

Two layouts are supported, recognised by their column *names*, never by
position, so reordered columns still work:

  northstar  one row per item and size
             required: Item Code, Size, Units Available, Cost USD
             optional: Description, Retail USD, Category
  harbor     one row per style, one column per size
             required: Style, Unit Cost USD, at least one size column
             optional: Product, Retail USD, Total Units

Anything else is rejected with a message naming the columns we looked for.
Title rows above the header and footer notes below the data are skipped and
reported, never parsed as items.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .normalize import Parsed, parse_item_code, parse_money, parse_quantity, parse_text

Layout = Literal["northstar", "harbor"]

HEADER_SCAN_ROWS = 30
MAX_DATA_ROWS = 50_000

NORTHSTAR_FIELDS: dict[str, str] = {
    "item code": "item_code",
    "description": "description",
    "size": "size",
    "units available": "quantity",
    "cost usd": "unit_cost",
    "retail usd": "retail",
    "category": "category",
}
NORTHSTAR_REQUIRED = ("item_code", "size", "quantity", "unit_cost")

HARBOR_FIELDS: dict[str, str] = {
    "style": "item_code",
    "product": "description",
    "unit cost usd": "unit_cost",
    "retail usd": "retail",
    "total units": "supplier_total",
}
HARBOR_REQUIRED = ("item_code", "unit_cost")

SIZE_TOKEN = re.compile(
    r"^(xxs|xs|s|m|l|xl|xxl|xxxl|[2-6]xl|os|one size|\d{1,3}(\.5)?|\d{1,2}[x/]\d{1,2})$",
    re.IGNORECASE,
)

FIELD_LABELS = {
    "item_code": "Item code",
    "description": "Description",
    "size": "Size",
    "quantity": "Pieces",
    "unit_cost": "Supplier cost",
    "retail": "Retail",
    "category": "Category",
    "supplier_total": "Supplier total",
}


class UnsupportedWorkbook(ValueError):
    """The file can't be read as one of the supported layouts."""


@dataclass(frozen=True)
class SourceCell:
    coordinate: str
    raw: Any  # JSON-safe copy of what the supplier's cell held


@dataclass
class SourceLine:
    line_id: str
    source_row: int
    item_code: Parsed[str]
    description: Parsed[str]
    size: Parsed[str]
    category: Parsed[str]
    quantity: Parsed[int]
    unit_cost: Parsed[Decimal]
    retail: Parsed[Decimal]
    cells: dict[str, SourceCell]
    row_warnings: list[str] = field(default_factory=list)


@dataclass
class ReadResult:
    layout: Layout
    sheet_name: str
    title: str | None
    supplier_name: str | None
    header_row: int
    columns: dict[str, str]  # field -> column letter (sizes as "size:M")
    lines: list[SourceLine]
    notices: list[str]


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip().rstrip(":").lower()


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if value == value and abs(value) != float("inf") else str(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (_dt.date, _dt.datetime, _dt.time)):
        return value.isoformat()
    return str(value)


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


@dataclass
class _HeaderMatch:
    layout: Layout
    row: int
    field_cols: dict[str, int]  # field -> 1-based column index
    size_cols: list[tuple[str, int]]  # (size label, column index)
    ignored: list[str]


def _match_header(row_idx: int, values: list[Any]) -> _HeaderMatch | None:
    headers = [(_norm_header(v), i + 1, v) for i, v in enumerate(values)]
    present = {h for h, _, _ in headers if h}

    for layout, fields, required in (
        ("northstar", NORTHSTAR_FIELDS, NORTHSTAR_REQUIRED),
        ("harbor", HARBOR_FIELDS, HARBOR_REQUIRED),
    ):
        mapped = {fields[h]: h for h in present if h in fields}
        if not all(r in mapped for r in required):
            continue

        field_cols: dict[str, int] = {}
        size_cols: list[tuple[str, int]] = []
        ignored: list[str] = []
        for norm, col, raw in headers:
            if not norm:
                continue
            if norm in fields:
                name = fields[norm]
                if name in field_cols:
                    raise UnsupportedWorkbook(
                        f"Column '{raw}' appears twice in the header row (row {row_idx}). "
                        "Remove the duplicate so we know which one to use."
                    )
                field_cols[name] = col
            elif layout == "harbor" and SIZE_TOKEN.fullmatch(norm):
                label = str(raw).strip().upper()
                if any(label == s for s, _ in size_cols):
                    raise UnsupportedWorkbook(
                        f"Size column '{raw}' appears twice in the header row (row {row_idx})."
                    )
                size_cols.append((label, col))
            else:
                ignored.append(str(raw).strip())

        if layout == "harbor" and not size_cols:
            continue
        return _HeaderMatch(layout, row_idx, field_cols, size_cols, ignored)
    return None


def _supplier_from_title(title: str | None) -> str | None:
    if not title:
        return None
    for sep in (" - ", " – ", " — ", ": "):
        if sep in title:
            return title.split(sep, 1)[0].strip() or None
    return title.strip() or None


def read_workbook(data: bytes, *, read_only: bool = True) -> ReadResult:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=read_only)
    except Exception as exc:  # openpyxl raises many types for non-xlsx input
        raise UnsupportedWorkbook(
            "This file couldn't be opened as an Excel workbook (.xlsx). "
            "Save it as .xlsx and upload again."
        ) from exc

    try:
        tried: list[str] = []
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            result = _read_sheet(ws)
            if result is not None:
                other = [
                    s.title for s in wb.worksheets
                    if s.title != ws.title and getattr(s, "sheet_state", "visible") == "visible"
                ]
                if other:
                    result.notices.append(
                        f"Read sheet '{ws.title}'. Other sheets were not read: {', '.join(other)}."
                    )
                return result
            tried.append(ws.title)
    finally:
        wb.close()

    raise UnsupportedWorkbook(
        "We couldn't find a supported header row"
        + (f" in sheet(s) {', '.join(repr(t) for t in tried)}" if tried else "")
        + ". Northstar sheets need the columns Item Code, Size, Units Available and Cost USD. "
        "Harbor sheets need Style, Unit Cost USD and size columns such as S, M, L."
    )


def _read_sheet(ws) -> ReadResult | None:
    rows = ws.iter_rows()
    preamble: list[str] = []
    match: _HeaderMatch | None = None

    for row_idx, row in enumerate(rows, start=1):
        values = [c.value for c in row]
        match = _match_header(row_idx, values)
        if match:
            break
        texts = [str(v).strip() for v in values if not _is_empty(v)]
        if texts:
            preamble.append(" ".join(texts))
        if row_idx >= HEADER_SCAN_ROWS:
            return None
    if match is None:
        return None

    title = preamble[0] if preamble else None
    columns = {f: get_column_letter(c) for f, c in match.field_cols.items()}
    columns.update({f"size:{s}": get_column_letter(c) for s, c in match.size_cols})

    notices: list[str] = []
    if match.ignored:
        notices.append(f"Ignored column(s) we don't use: {', '.join(match.ignored)}.")

    reader = _read_northstar_rows if match.layout == "northstar" else _read_harbor_rows
    lines = reader(rows, match, notices)

    return ReadResult(
        layout=match.layout,
        sheet_name=ws.title,
        title=title,
        supplier_name=_supplier_from_title(title),
        header_row=match.row,
        columns=columns,
        lines=lines,
        notices=notices,
    )


def _cell_at(row: tuple, col: int):
    return row[col - 1] if col - 1 < len(row) else None


def _grab(row: tuple, col: int | None) -> tuple[Any, str | None, str | None]:
    """(value, number_format, coordinate) for a 1-based column, tolerating short rows."""
    if col is None:
        return None, None, None
    cell = _cell_at(row, col)
    if cell is None:
        return None, None, None
    coord = getattr(cell, "coordinate", None)
    return cell.value, getattr(cell, "number_format", None), coord


def _row_number(row: tuple, fallback: int) -> int:
    for cell in row:
        r = getattr(cell, "row", None)
        if r:
            return r
    return fallback


def _skip_or_note(row_idx: int, row: tuple, value_cols: list[int], notices: list[str]) -> bool:
    """True if this row holds no item data. Text-only rows are reported."""
    if any(not _is_empty(_grab(row, c)[0]) for c in value_cols):
        return False
    texts = [str(c.value).strip() for c in row if c is not None and not _is_empty(c.value)]
    if texts:
        snippet = " ".join(texts)
        snippet = snippet if len(snippet) <= 80 else snippet[:77] + "..."
        notices.append(f"Row {row_idx} skipped: no item data (\"{snippet}\").")
    return True


def _read_northstar_rows(rows, match: _HeaderMatch, notices: list[str]) -> list[SourceLine]:
    fc = match.field_cols
    value_cols = [fc[f] for f in ("description", "size", "quantity", "unit_cost", "retail", "category") if f in fc]
    lines: list[SourceLine] = []

    for offset, row in enumerate(rows, start=match.row + 1):
        row_idx = _row_number(row, offset)
        if _skip_or_note(row_idx, row, value_cols, notices):
            continue
        if len(lines) >= MAX_DATA_ROWS:
            raise UnsupportedWorkbook(f"More than {MAX_DATA_ROWS:,} item rows; split the file.")

        cells: dict[str, SourceCell] = {}
        raw: dict[str, tuple[Any, str | None]] = {}
        for name, col in fc.items():
            value, fmt, coord = _grab(row, col)
            raw[name] = (value, fmt)
            cells[name] = SourceCell(coord or f"{get_column_letter(col)}{row_idx}", _json_safe(value))

        lines.append(SourceLine(
            line_id=f"R{row_idx}",
            source_row=row_idx,
            item_code=parse_item_code(*raw["item_code"]),
            description=parse_text(raw.get("description", (None, None))[0]),
            size=parse_text(raw["size"][0]),
            category=parse_text(raw.get("category", (None, None))[0]),
            quantity=parse_quantity(raw["quantity"][0]),
            unit_cost=parse_money(raw["unit_cost"][0]),
            retail=parse_money(raw.get("retail", (None, None))[0]),
            cells=cells,
        ))
    return lines


def _read_harbor_rows(rows, match: _HeaderMatch, notices: list[str]) -> list[SourceLine]:
    fc = match.field_cols
    value_cols = [c for _, c in match.size_cols] + [
        fc[f] for f in ("description", "unit_cost", "retail", "supplier_total") if f in fc
    ]
    lines: list[SourceLine] = []

    for offset, row in enumerate(rows, start=match.row + 1):
        row_idx = _row_number(row, offset)
        if _skip_or_note(row_idx, row, value_cols, notices):
            continue

        shared_cells: dict[str, SourceCell] = {}
        shared_raw: dict[str, tuple[Any, str | None]] = {}
        for name, col in fc.items():
            value, fmt, coord = _grab(row, col)
            shared_raw[name] = (value, fmt)
            shared_cells[name] = SourceCell(coord or f"{get_column_letter(col)}{row_idx}", _json_safe(value))

        item_code = parse_item_code(*shared_raw["item_code"])
        description = parse_text(shared_raw.get("description", (None, None))[0])
        unit_cost = parse_money(shared_raw["unit_cost"][0])
        retail = parse_money(shared_raw.get("retail", (None, None))[0])

        row_lines: list[SourceLine] = []
        blank_sizes: list[str] = []
        for size_label, col in match.size_cols:
            value, _, coord = _grab(row, col)
            if _is_empty(value):
                blank_sizes.append(size_label)
                continue
            cells = dict(shared_cells)
            cells.pop("supplier_total", None)
            cells["quantity"] = SourceCell(coord or f"{get_column_letter(col)}{row_idx}", _json_safe(value))
            row_lines.append(SourceLine(
                line_id=f"R{row_idx}-{size_label}",
                source_row=row_idx,
                item_code=item_code,
                description=description,
                size=Parsed(size_label, "ok"),
                category=Parsed(None, "missing"),
                quantity=parse_quantity(value),
                unit_cost=unit_cost,
                retail=retail,
                cells=cells,
            ))

        if len(lines) + len(row_lines) > MAX_DATA_ROWS:
            raise UnsupportedWorkbook(f"More than {MAX_DATA_ROWS:,} item lines; split the file.")

        if blank_sizes and row_lines:
            notices.append(
                f"Row {row_idx}: size(s) {', '.join(blank_sizes)} left blank by the supplier, treated as not offered."
            )
        elif blank_sizes and not row_lines:
            notices.append(f"Row {row_idx}: every size is blank, so no lines were created for this row.")

        # The supplier's own total is a claim to check, never a value we use.
        total = parse_quantity(shared_raw.get("supplier_total", (None, None))[0])
        if total.usable and row_lines:
            usable = sum(
                ln.quantity.value for ln in row_lines
                if ln.quantity.usable and ln.quantity.value is not None and ln.quantity.value > 0
            )
            if total.value != usable:
                warning = (
                    f"Supplier's Total Units for this row says {total.value:,}, "
                    f"but the usable pieces across sizes add up to {usable:,}."
                )
                for ln in row_lines:
                    ln.row_warnings.append(warning)

        lines.extend(row_lines)
    return lines
